import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_code_block(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F4F5F7")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    # Border
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for b_name in ['top', 'bottom', 'right']:
        b = OxmlElement(f'w:{b_name}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), 'E2E8F0')
        tcBorders.append(b)
    left_b = OxmlElement('w:left')
    left_b.set(qn('w:val'), 'single')
    left_b.set(qn('w:sz'), '24') # thick left border
    left_b.set(qn('w:space'), '0')
    left_b.set(qn('w:color'), '3B82F6')
    tcBorders.append(left_b)
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph()

def style_table(table, col_widths, headers, data):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Header row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1E3A8A") # Navy Blue
        set_cell_margins(hdr_cells[i], top=140, bottom=140, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(255, 255, 255)
            
    # Data rows
    for r_idx, row_data in enumerate(data):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=100, bottom=100, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            for run in p.runs:
                run.font.name = 'Calibri'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(51, 65, 85)

    # Set col widths
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = Inches(width)

def build_docx(filename):
    doc = Document()
    
    # Set page margins (1 inch)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    trun = title_p.add_run("BÁO CÁO THIẾT KẾ CÁC LUỒNG HỆ THỐNG")
    trun.font.name = 'Calibri'
    trun.font.size = Pt(22)
    trun.font.bold = True
    trun.font.color.rgb = RGBColor(30, 58, 138)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(18)
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    srun = sub_p.add_run("Dự án: LabCast — Phòng thực hành DSA & SQL có animation đồng bộ mạng LAN\nHọc phần: Lập trình mạng (Network Programming)")
    srun.font.name = 'Calibri'
    srun.font.size = Pt(12)
    srun.font.italic = True
    srun.font.color.rgb = RGBColor(71, 85, 105)

    # Divider line
    div_p = doc.add_paragraph()
    div_p.paragraph_format.space_after = Pt(12)
    p_run = div_p.add_run("_________________________________________________________________________________")
    p_run.font.color.rgb = RGBColor(203, 213, 225)

    # 1. Giới thiệu tổng quan
    h1 = doc.add_heading("1. TỔNG QUAN KIẾN TRÚC MẠNG & CÁC KÊNH TRUYỀN THÔNG", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    
    p = doc.add_paragraph(
        "Hệ thống LabCast phục vụ mô hình phòng máy thực hành mạng dây (tối đa 50 máy học viên kết nối chung switch với máy giáo viên). "
        "Nhằm đáp ứng yêu cầu truyền thông thời gian thực, tiết kiệm băng thông và đảm bảo độ tin cậy cao, kiến trúc phân bổ thành 3 kênh giao tiếp chính:"
    )
    p.paragraph_format.line_spacing = 1.2
    
    # Table channels
    table1 = doc.add_table(rows=5, cols=4)
    headers1 = ["Kênh / Cổng", "Giao thức", "Cơ chế chính", "Mục đích sử dụng"]
    widths1 = [1.2, 1.1, 1.6, 2.6]
    data1 = [
        ["Cổng 7000", "TCP", "Protocol LCP/1.0 (nhị phân / JSON)", "Phiên làm việc bền vững: đăng nhập, bài học, nhật ký phím gõ, chạy thử, nộp bài, câu hỏi nhanh, trợ giúp"],
        ["239.255.70.1:7001", "UDP Multicast", "Multicast tin cậy có NACK & Keyframe", "Chiếu animation giảng bài đồng loạt cho cả phòng máy; tự động lùi về TCP nếu bị chặn"],
        ["Cổng 7002", "UDP Broadcast", "Gói tin DISCOVER / REPLY", "Ứng dụng máy trạm tự động tìm kiếm địa chỉ Server trong mạng LAN khi khởi động"],
        ["Cổng 7000 (TLS)", "TLS trên TCP", "Mã hoá TLS tuỳ chọn", "Bảo vệ phiên làm việc và bảo mật xác thực (tùy chọn qua file cấu hình)"]
    ]
    style_table(table1, widths1, headers1, data1)
    doc.add_paragraph()

    # 2. Các luồng giao tiếp mạng và nghiệp vụ
    h2 = doc.add_heading("2. CÁC LUỒNG GIAO TIẾP MẠNG & NGHIỆP VỤ (COMMUNICATION FLOWS)", level=1)
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(8)

    # 2.1 Luồng vào lớp
    h2_1 = doc.add_heading("2.1. Luồng Tự động tìm Server, Bắt tay và Gia nhập lớp (Discovery & Handshake)", level=2)
    doc.add_paragraph(
        "Mục tiêu của luồng là giúp máy trạm tự cấu hình địa chỉ mạng, chứng thực danh tính, đồng bộ đồng hồ với độ trễ thấp và gia nhập nhóm chiếu multicast."
    )
    
    seq1_code = """App học viên                                  Server
  |== DISCOVER (UDP broadcast :7002) ========>|
  |<= DISCOVER_REPLY(lớp, cổng 7000, TLS) ----|
  |                                           |
  |---- HELLO(ver=1.0) (TCP :7000) ---------->|
  |<--- HELLO_ACK(serverTime, classId) -------|
  |---- AUTH(username, password) ------------>| (Xác thực PBKDF2)
  |<--- AUTH_OK(token, role, classId) --------|
  |                                           |
  |---- TIME_SYNC_REQ(t1) x 5 --------------->| (Đồng bộ đồng hồ ĐG3)
  |<--- TIME_SYNC_RESP(t2, t3) ---------------| (Lấy mẫu RTT nhỏ nhất)
  |                                           |
  |---- CLASS_JOIN -------------------------->| -> Cập nhật Sơ đồ lớp (CLASS_STATE)
  |     (Tham gia 239.255.70.1 trên card LAN) |
  |---- LESSON_LIST ------------------------->|
  |<--- LESSON_DATA(danh sách & khoá) --------|"""
    add_code_block(doc, seq1_code)

    doc.add_paragraph(
        "Đặc điểm kỹ thuật nổi bật:\n"
        "• Lựa chọn đúng card mạng vật lý: Sau khi bắt tay TCP, Client truy vấn socket.getLocalAddress() để xác định đúng NetworkInterface vật lý kết nối vào LAN, tránh nghe nhầm card ảo (VMware, VirtualBox, WSL).\n"
        "• Đồng bộ đồng hồ kiểu NTP (Đóng góp ĐG3): Đo 5 mẫu liên tiếp và chọn mẫu RTT nhỏ nhất để tính toán serverClockOffset = ((t2 - t1) + (t3 - t4)) / 2, bù trôi đồng hồ định kỳ."
    )

    # 2.2 Luồng chiếu animation
    h2_2 = doc.add_heading("2.2. Luồng Chiếu bài giảng Animation qua UDP Multicast (Broadcasting Flow)", level=2)
    doc.add_paragraph(
        "Thay vì truyền stream video màn hình tiêu tốn băng thông lớn (hàng chục Mbps), LabCast phát chuỗi sự kiện hoạt hoạ (trace events) qua multicast và đồng bộ phát theo thời gian thực."
    )

    seq2_code = """App giáo viên        Server                                 Các máy học viên
  |-CAST_CONTROL----->|  PLAY bước 37, tốc độ 1x                  |
  |  (TCP)            |== CAST_DATA x k (Multicast) =============>| (Trace nén Deflate <= 1387B)
  |                   |== CAST_CONTROL(PLAY, 37, at=now+300ms) ==>| (Phát đồng loạt tại t=now+300ms)
  |                   |== CAST_KEYFRAME (định kỳ mỗi 2s) ========>| (Máy vào muộn bắt kịp)
  |                   |                                           |
  |                   |<-------- CAST_NACK(từ seq, đến seq) (TCP)-| (Khi phát hiện khuyết gói)
  |                   |--------- CAST_REPAIR (TCP) -------------->| (Gửi bù gói thiếu qua TCP)"""
    add_code_block(doc, seq2_code)

    doc.add_paragraph(
        "Cơ chế tin cậy trên nền UDP:\n"
        "• Tránh phân mảnh IP (IP Fragmentation): Mỗi datagram multicast giới hạn tối đa 1400 bytes (13 bytes LCP Frame header + tối đa 1387 bytes payload).\n"
        "• Sửa lỗi mất gói: Máy trạm duy trì bộ đệm và phát hiện khuyết SEQ; gửi CAST_NACK qua kênh TCP tin cậy để nhận gói sửa CAST_REPAIR.\n"
        "• Hẹn giờ đồng bộ: Lệnh phát mang mốc thời gian tuyệt đối at = Clock.serverNow() + 300ms, giúp toàn bộ màn hình chuyển cảnh lệch nhau không quá 50ms (p95)."
    )

    # 2.3 Luồng làm bài và đồng bộ dòng code
    h2_3 = doc.add_heading("2.3. Luồng Gõ code, Chạy thử và Nộp bài (Reliable Coding & Evaluation)", level=2)
    doc.add_paragraph(
        "Giải quyết bài toán máy phòng lab hay treo, tắt nguồn đột ngột và chống gian lận nộp bài."
    )

    seq3_code = """App học viên                                  Server                     Docker Sandbox
  |---- LESSON_FETCH(id) -------------------->|                                |
  |<--- LESSON_DATA(đề, khung code) ----------|                                |
  |                                           |                                |
  |  (Gõ code: gộp phím 100ms, ghi ổ đĩa)     |                                |
  |---- CODE_DELTA(seq = n, thay đổi) ------->| (Lưu database journal)         |
  |<--- CODE_ACK(lastSeq = n) ----------------|                                |
  |                                           |                                |
  |---- RUN_REQ(codeSeq = n) ---------------->| ──► Đẩy vào hàng đợi Sandbox ─►│
  |<--- RUN_QUEUED(vị trí hàng đợi) ----------|                                │
  |<--- TRACE_CHUNK x k (Stream animation) ---|<── Đọc trace stdout từ Docker ─┘
  |<--- RUN_RESULT(test xem: PASS) -----------|                                
  |                                           |                                
  |---- SUBMIT_REQ(codeSeq = n) ------------->| ──► Chạy toàn bộ test ẩn ─────►│
  |<--- SUBMIT_RESULT(điểm, test, cooldown) --|<── Kết quả chấm ---------------┘"""
    add_code_block(doc, seq3_code)

    doc.add_paragraph(
        "Điểm kỹ thuật quan trọng:\n"
        "• Dòng code tin cậy (Đóng góp ĐG2): Client lưu nhật ký cục bộ dạng append-only trước khi gửi mạng; Server lưu lịch sử delta; khi mất điện hoặc đổi máy chỉ cần gửi RESUME_REQ để đồng bộ lại.\n"
        "• An toàn với Docker Sandbox: Code chỉ thực thi trong container dùng một lần với --network none, giới hạn RAM 256MB, CPU 1 core, ngắt tiến trình bằng RLIMIT_CPU."
    )

    # 2.4 Luồng câu hỏi nhanh
    h2_4 = doc.add_heading("2.4. Luồng Câu hỏi nhanh tính giờ công bằng (Fair Quiz Flow)", level=2)
    doc.add_paragraph(
        "Giáo viên phát câu hỏi trắc nghiệm tương tác nhanh (ví dụ: 'Cặp phần tử nào đổi chỗ tiếp theo? - 10s'). Hệ thống áp dụng cơ chế bù độ trễ mạng ĐG4:"
    )

    seq4_code = """App giáo viên        Server                                 Các máy học viên
  |-QUIZ_OPEN-------->|--------- QUIZ_OPEN(đề, opensAt, closesAt) ->| (Gửi trước >= 500ms)
  |                   |          (opensAt = now + 500ms, 10s)       |
  |                   |                                             |
  |                   |<-------- QUIZ_ANSWER(đáp án, clientTime) ---|
  |                   |          [Tính điểm: Server trừ RTT/2 có trần]
  |                   |--------- QUIZ_CLOSE ----------------------->|
  |<-QUIZ_RESULT------|--------- QUIZ_RESULT ---------------------->|"""
    add_code_block(doc, seq4_code)

    # 2.5 Luồng trợ giúp
    h2_5 = doc.add_heading("2.5. Luồng Trợ giúp, Giơ tay và Phản chiếu Code (Teacher Assist & Mirroring)", level=2)
    seq5_code = """App học viên              Server                       App giáo viên
  |-- HAND_RAISE --------->|-- CLASS_STATE(✋) --------->|
  |                        |<-- MIRROR_SUBSCRIBE(hv) ----|
  |                        |-- MIRROR_SNAPSHOT --------->| (Code đầy đủ + seq)
  |-- CODE_DELTA --------->|-- MIRROR_DELTA ------------>| (Phản chiếu <= 200ms)
  |                        |<-- COMMENT_ADD(dòng, nội dung)|
  |<-- COMMENT ------------|                             |
  |-- HAND_LOWER --------->|<-- MIRROR_UNSUBSCRIBE ------|"""
    add_code_block(doc, seq5_code)

    # 3. Kiến trúc Concurrency
    h3 = doc.add_heading("3. KIẾN TRÚC LUỒNG TIẾN TRÌNH & ĐỒNG THỜI (CONCURRENCY & THREADING)", level=1)
    h3.paragraph_format.space_before = Pt(16)
    h3.paragraph_format.space_after = Pt(8)

    doc.add_paragraph(
        "Để đáp ứng đồng thời 50 máy trạm liên tục gửi phím gõ và chạy code mà không làm nghẽn hệ thống, mô hình đa luồng được phân tách nhiệm vụ rõ rệt giữa Server và Client:"
    )

    h3_1 = doc.add_heading("3.1. Phía Server (Java Server)", level=2)
    server_threads = [
        ["Accept Thread", "1 luồng", "ServerSocket.accept() đón kết nối mới, khởi tạo luồng đọc/ghi"],
        ["Client Read Thread", "1 luồng / client (tối đa 50)", "Đọc tuần tự frame từ Socket bằng vòng lặp readFully(), kiểm tra MAGIC, giải nén và chuyển sang Worker Pool"],
        ["Client Write Thread", "1 luồng / client (tối đa 50)", "Lấy gói tin từ hàng đợi gửi (PriorityBlockingQueue), thực hiện ghi nhị phân ra Socket"],
        ["Business Worker Pool", "Cố định (2 x Cores)", "ThreadPoolExecutor xử lý logic nghiệp vụ: tính toán điểm, cập nhật sơ đồ lớp, ghi database SQLite"],
        ["Multicast Cast Loop", "1 luồng", "Gửi datagram ra nhóm multicast; tự động giãn nhịp (Throttling <= 2 MB/s) để chống tràn bộ đệm máy trạm"],
        ["Runner Sandbox Pool", "Giới hạn (Cores - 1)", "Quản lý tiến trình Docker; mỗi học viên chỉ được xếp 1 lượt chạy thử đang chờ"]
    ]
    t_server = doc.add_table(rows=7, cols=3)
    style_table(t_server, [1.6, 1.4, 3.5], ["Thành phần luồng", "Số lượng luồng", "Trách nhiệm chính"], server_threads)
    doc.add_paragraph()

    doc.add_paragraph(
        "Cơ chế Backpressure trên hàng đợi ghi (Write Queue) của Server:\n"
        "Mỗi kết nối có hàng đợi 256 phần tử có phân loại mức độ ưu tiên:\n"
        "• CRITICAL (AUTH, QUIZ, SUBMIT_RESULT, ERROR, BYE): Luôn ưu tiên chen lên đầu, không bao giờ bị huỷ bỏ.\n"
        "• NORMAL (CODE_ACK, LESSON_DATA...): Chờ tối đa 5 giây; nếu nghẽn mạng thì đánh dấu kết nối SUSPECT.\n"
        "• CONFLATABLE (MIRROR_DELTA, CLASS_STATE): Khi hàng đợi đầy, huỷ các bản delta cũ đang chờ và thay bằng 1 bản chụp (SNAPSHOT) trạng thái mới nhất.\n"
        "• DROPPABLE (Gói chiếu TCP dự phòng): Bỏ gói cũ nhất, đợi keyframe kế tiếp tự động sửa trạng thái."
    )

    h3_2 = doc.add_heading("3.2. Phía Client (JavaFX Desktop App)", level=2)
    client_threads = [
        ["JavaFX Application Thread", "1 luồng UI chính", "Xử lý tương tác giao diện người dùng, chạy AnimationTimer (60 FPS) tính frame theo Clock.serverNow(). Tuyệt đối không thực hiện I/O mạng."],
        ["TCP Read Thread", "1 luồng", "Lắng nghe frame từ Server qua Socket; giải mã và chuyển dữ liệu sang UI bằng Platform.runLater()."],
        ["TCP Send Worker", "1 luồng", "Gửi gói tin theo thứ tự lên Server qua Socket outstream."],
        ["Multicast Listener Thread", "1 luồng", "Lắng nghe MulticastSocket trên cổng 7001; kiểm tra SEQ, phát hiện mất gói và gọi hàm ráp trace."],
        ["Local Journal Thread", "1 luồng đơn", "SingleThreadExecutor ghi phím gõ tuần tự xuống file nhật ký đĩa cứng cục bộ trước khi truyền qua mạng."]
    ]
    t_client = doc.add_table(rows=6, cols=3)
    style_table(t_client, [1.8, 1.4, 3.3], ["Thành phần luồng", "Số lượng luồng", "Trách nhiệm chính"], client_threads)
    doc.add_paragraph()

    # 4. Bảng tổng hợp Protocol Frames
    h4 = doc.add_heading("4. BẢNG TỔNG HỢP CẤU TRÚC FRAME PROTOCOL LCP/1.0", level=1)
    h4.paragraph_format.space_before = Pt(16)
    h4.paragraph_format.space_after = Pt(8)

    doc.add_paragraph(
        "Định dạng Frame nhị phân đồng nhất giữa TCP và UDP (tổng 13 bytes header):"
    )

    frame_ascii = """ 0        2      3      4      5              9             13
 +--------+------+------+------+--------------+--------------+----------------+
 | MAGIC  | VER  | TYPE | FLAGS|     SEQ      |     LEN      |    PAYLOAD     |
 | 'L''C' | 1B   | 1B   | 1B   |    4B (BE)   |    4B (BE)   |   LEN bytes    |
 +--------+------+------+------+--------------+--------------+----------------+"""
    add_code_block(doc, frame_ascii)

    frame_fields = [
        ["MAGIC (2B)", "0x4C, 0x43 ('LC')", "Nhận diện giao thức LabCast, loại bỏ client lạ hoặc stream sai ngay từ 2 byte đầu."],
        ["VER (1B)", "0x01", "Phiên bản giao thức (LCP/1.0)."],
        ["TYPE (1B)", "0x01 - 0x50", "Mã định danh loại thông điệp (AUTH, CODE_DELTA, CAST_DATA, v.v.)."],
        ["FLAGS (1B)", "Bitmask 8-bit", "Bit 0: COMPRESSED (Deflate); Bit 1: JSON; Bit 2: REQUIRES_ACK; Bit 3: VIA_TCP."],
        ["SEQ (4B)", "Big-Endian Integer", "Số thứ tự tăng dần theo kết nối hoặc thứ tự luồng multicast."],
        ["LEN (4B)", "Big-Endian Integer", "Độ dài dữ liệu Payload (tối đa 8MB). Luôn kiểm tra kích thước trước khi cấp phát RAM."],
        ["PAYLOAD", "LEN bytes", "Nội dung thông điệp (nhị phân nhúng hoặc JSON UTF-8 tuỳ cờ FLAGS)."]
    ]
    t_frame = doc.add_table(rows=8, cols=3)
    style_table(t_frame, [1.4, 1.6, 3.5], ["Trường Header", "Định dạng", "Ý nghĩa kỹ thuật"], frame_fields)

    # Conclusion
    doc.add_paragraph()
    h5 = doc.add_heading("5. KẾT LUẬN", level=1)
    h5.paragraph_format.space_before = Pt(14)
    h5.paragraph_format.space_after = Pt(6)
    p_concl = doc.add_paragraph(
        "Thiết kế các luồng giao tiếp và luồng xử lý đồng thời trong LabCast tập trung trực tiếp vào bản chất bài toán mạng: "
        "kết hợp tính tin cậy của TCP cho phiên làm việc với ưu thế băng thông của UDP Multicast cho trình chiếu; "
        "bù trừ độ trễ đồng hồ phân tán theo chuẩn NTP; và kiến trúc đa luồng phân tách độc lập ngăn chặn hiện tượng tắc nghẽn giao diện. "
        "Tài liệu này là cơ sở vững chắc cho phần thiết kế truyền thông (Network Communication Design) và kiến trúc hệ thống trong báo cáo bài tập lớn."
    )
    p_concl.paragraph_format.line_spacing = 1.2

    # Save
    doc.save(filename)
    print(f"File created successfully: {filename}")

if __name__ == "__main__":
    out_dir = r"d:\Download\mai xoa\ExamNet\report"
    os.makedirs(out_dir, exist_ok=True)
    out_file = os.path.join(out_dir, "Bao_Cao_Cac_Luong_He_Thong_LabCast.docx")
    build_docx(out_file)
